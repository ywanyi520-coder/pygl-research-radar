#!/usr/bin/env python3
"""Targeted, processed-data-only audit of non-CPTAC human PYGL Ser15 cohorts.

This script is intentionally independent of the frozen CPTAC/PDC route. It
queries only the nine explicitly requested ProteomeXchange datasets, downloads
public RESULT/SEARCH/OTHER files (never RAW or PEAK files), and stops the
quantitative search after the first sequence-confirmed patient-level S15 hit.
"""
from __future__ import annotations

import csv
import gzip
import io
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any, Iterable

import requests

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
DOWNLOADS = ROOT / "remote_audit_data"
MANIFEST = RESULTS / "NonCPTAC_S15_file_manifest.tsv"
INVENTORY = RESULTS / "NonCPTAC_S15_dataset_inventory.tsv"
EVIDENCE = RESULTS / "NonCPTAC_S15_candidate_evidence.tsv"
STATUS = RESULTS / "NonCPTAC_S15_audit_status.txt"
PX_HTML = "https://proteomecentral.proteomexchange.org/cgi/GetDataset?ID={accession}"
IPROX_API = "https://www.iprox.cn/proxi/datasets/{accession}"
JPOST_API = "https://repository.jpostdb.org/proxi/datasets/{accession}"

# Public cohort facts are drawn from the primary papers/dataset records cited
# in the accompanying audit report. sample_n refers to biological tissues or
# treated patient aliquots, not mass-spectrometry injections.
STUDIES: list[dict[str, str]] = [
    dict(accession="PXD014997", tier="1", disease="AML",
         patient_n="41", sample_n="41", tissue="Diagnostic primary AML cells; bone marrow and/or peripheral blood",
         clinical_endpoint="5-year relapse vs relapse-free status",
         total_proteome_expected="YES", mapping_expected="YES (published clinical cohort metadata)"),
    dict(accession="PXD053618", tier="1", disease="FLT3-mutated AML",
         patient_n="47", sample_n="71", tissue="Diagnosis peripheral blood (37) and bone marrow (34); repeated patients",
         clinical_endpoint="Midostaurin + intensive chemotherapy response; DFS/EFS and long-term outcome; validation cohorts",
         total_proteome_expected="UNKNOWN", mapping_expected="RESTRICTED (ethics approval/data-sharing agreement stated)"),
    dict(accession="PXD032110", tier="1", disease="AML (TCGA-LAML)",
         patient_n="44", sample_n="50 (44 AML + 6 healthy marrow controls)", tissue="TCGA AML bone-marrow specimens and healthy bone marrow controls",
         clinical_endpoint="Overall survival, risk, mutation, clinical outcome via exact TCGA case-ID mapping",
         total_proteome_expected="YES", mapping_expected="YES IF exact TCGA case IDs are present"),
    dict(accession="PXD044246", tier="1", disease="Colorectal cancer",
         patient_n="53", sample_n="54 phosphoproteome samples (within a 387-patient cohort)", tissue="Primary colorectal tumor biopsies",
         clinical_endpoint="Post-surgery relapse / relapse-free survival (availability in phospho subset to verify)",
         total_proteome_expected="YES (larger cohort)", mapping_expected="UNKNOWN (patient link to phospho subset to verify)"),
    dict(accession="PXD038081", tier="1", disease="APC-mutant colorectal cancer",
         patient_n="69", sample_n="138 paired tumor/adjacent tissues", tissue="Colon tumor and paired adjacent tissue; APC-mutant subgroup n=35",
         clinical_endpoint="Overall survival / prognosis; chemotherapy response reported",
         total_proteome_expected="YES", mapping_expected="YES IF sample identifiers align to clinical supplement"),
    dict(accession="PXD048988", tier="1", disease="Recurrent colorectal liver metastasis",
         patient_n="24", sample_n="96 (metastasis and normal liver at initial and recurrent resections, per dataset description)",
         tissue="Paired initial/recurrent liver metastases and normal liver",
         clinical_endpoint="Paired initial-to-recurrence change; prognosis in publication",
         total_proteome_expected="YES", mapping_expected="YES IF patient-paired sample IDs are retained"),
    dict(accession="PXD005173", tier="1", disease="Colorectal cancer",
         patient_n="20", sample_n="40", tissue="Tumor and matched tumor-adjacent normal colon tissue",
         clinical_endpoint="Paired tumor vs adjacent normal (human disease relevance only)",
         total_proteome_expected="NO/UNKNOWN (phosphoproteome submission)", mapping_expected="YES IF pair labels are retained"),
    dict(accession="PXD034355", tier="2", disease="HER2-negative breast cancer",
         patient_n="130", sample_n="130", tissue="Pretreatment breast tumor biopsies from neoadjuvant cohort",
         clinical_endpoint="Pathologic complete response / paclitaxel sensitivity",
         total_proteome_expected="UNKNOWN (phosphoproteomic screen)", mapping_expected="YES IF clinical sample IDs align to pCR supplement"),
    dict(accession="PXD017660", tier="2", disease="AML",
         patient_n="20", sample_n="40 patient aliquots (paired DMSO vs selinexor; 4 cell lines excluded)",
         tissue="Ex-vivo primary AML blasts from blood or bone marrow",
         clinical_endpoint="Ex-vivo selinexor EC50/response; not an in-vivo clinical endpoint",
         total_proteome_expected="UNKNOWN", mapping_expected="YES IF patient IDs align to EC50 supplement"),
]

INVENTORY_FIELDS = [
    "dataset", "disease", "patient_n", "sample_n", "tissue", "clinical_endpoint",
    "processed_phosphosite_available", "total_proteome_available",
    "patient_mapping_available", "raw_only", "status",
]
MANIFEST_FIELDS = ["dataset", "repository_accession", "filename", "category", "size_bytes", "download_status"]
EVIDENCE_FIELDS = [
    "dataset", "repository_accession", "filename", "table_member", "row_number",
    "hit_type", "protein_identifiers", "site_identifiers", "peptide_sequence",
    "sequence_confirmed", "quantitative_sample_n", "observed_patient_n",
    "patient_ids_recovered", "site_localization", "false_positive_excluded",
]

SESSION = requests.Session()
SESSION.headers.update({"User-Agent": "nonCPTAC-PYGL-S15-audit/1.0 (public processed-data review)"})


def write_tsv(path: Path, fields: list[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: "" if row.get(key) is None else row.get(key, "") for key in fields})


def get_json(url: str) -> Any:
    response = SESSION.get(url, timeout=(15, 90))
    response.raise_for_status()
    return response.json()


def nested_strings(obj: Any) -> list[str]:
    if isinstance(obj, str):
        return [obj]
    if isinstance(obj, dict):
        return [s for value in obj.values() for s in nested_strings(value)]
    if isinstance(obj, list):
        return [s for value in obj for s in nested_strings(value)]
    return []


def resolve_repository_accession(pxd: str) -> tuple[str, str]:
    """Resolve a PXD accession to its native public repository accession."""
    # These are explicit, source-documented cross-repository mappings.
    known = {"PXD032110": "MSV000089012", "PXD005173": "JPST000210"}
    if pxd in known:
        return known[pxd], "known ProteomeXchange cross-reference"
    for url in (IPROX_API.format(accession=pxd), JPOST_API.format(accession=pxd)):
        try:
            payload = get_json(url)
            content = "\n".join(nested_strings(payload)) + "\n" + json.dumps(payload)
            match = re.search(r"\b(?:IPX\d{10}|JPST\d{6}|MSV\d{9})\b", content)
            if match:
                return match.group(0), url
        except Exception:
            pass
    try:
        page = SESSION.get(PX_HTML.format(accession=pxd), timeout=(15, 90))
        page.raise_for_status()
        match = re.search(r"\b(?:IPX\d{10}|JPST\d{6}|MSV\d{9})\b", page.text)
        if match:
            return match.group(0), PX_HTML.format(accession=pxd)
    except Exception:
        pass
    return pxd, "ProteomeXchange accession (native host not separately resolved)"


def list_processed_files(client: Any, pxd: str) -> tuple[str, str, list[dict[str, Any]], str]:
    native, source = resolve_repository_accession(pxd)
    errors: list[str] = []
    for accession in dict.fromkeys((native, pxd)):
        try:
            processed = client.get_all_category_file_list(accession, ["RESULT", "SEARCH", "OTHER"])
            raw_files = client.get_all_category_file_list(accession, ["RAW"])
            files = list(processed or []) + list(raw_files or [])
            if files:
                return native, source, files, "OK"
            errors.append(f"{accession}: no files listed")
        except Exception as exc:
            errors.append(f"{accession}: {type(exc).__name__}: {exc}")
    return native, source, [], "; ".join(errors)[:800]


def file_fields(record: dict[str, Any]) -> tuple[str, str, int]:
    name = str(record.get("fileName") or record.get("name") or record.get("filename") or "")
    category = record.get("fileCategory", "")
    if isinstance(category, dict):
        category = category.get("value", "")
    category = str(category).upper()
    size = record.get("fileSizeBytes", record.get("fileSize", record.get("size", 0)))
    try:
        size = int(size or 0)
    except (TypeError, ValueError):
        size = 0
    return name, category, size


TABLE_EXTENSIONS = (".txt", ".tsv", ".csv", ".xlsx", ".xls", ".txt.gz", ".tsv.gz", ".csv.gz", ".zip")
NAME_HINTS = re.compile(r"phosph|site|evidence|maxquant|protein|quant|matrix|tmt|report|result|table|supp|clinical|sample|peptide|summary|output|group|diann", re.I)
RAW_EXTENSIONS = (".raw", ".wiff", ".mzml", ".mzxml", ".d", ".mgf", ".baf", ".ibd", ".tdf", ".fid")
MAX_FILE_BYTES = 1_200_000_000
MAX_TOTAL_BYTES_PER_DATASET = 3_500_000_000


def choose_downloads(files: list[dict[str, Any]]) -> list[dict[str, Any]]:
    eligible: list[dict[str, Any]] = []
    for record in files:
        name, category, size = file_fields(record)
        lower = name.lower()
        if not name or category in {"RAW", "PEAK", "FASTA"}:
            continue
        if lower.endswith(RAW_EXTENSIONS) or not lower.endswith(TABLE_EXTENSIONS):
            continue
        if size > MAX_FILE_BYTES:
            continue
        if category in {"RESULT", "SEARCH"} or NAME_HINTS.search(name):
            eligible.append(record)
    return eligible


def download_selected(client: Any, accession: str, pxd: str, records: list[dict[str, Any]]) -> tuple[list[Path], list[dict[str, Any]]]:
    out = DOWNLOADS / pxd
    out.mkdir(parents=True, exist_ok=True)
    names: list[str] = []
    total = 0
    accepted: list[dict[str, Any]] = []
    seen: set[str] = set()
    for record in records:
        name, category, size = file_fields(record)
        if name in seen:
            continue
        if total + size > MAX_TOTAL_BYTES_PER_DATASET:
            continue
        seen.add(name)
        names.append(name)
        total += size
        accepted.append(record)
    if names:
        client.download_files_by_list(
            accession=accession,
            file_names=names,
            output_folder=str(out),
            skip_if_downloaded_already=True,
            protocol="ftp",
            aspera_maximum_bandwidth="100M",
            checksum_check=False,
            parallel_files=2,
            flatten=True,
        )
    paths = [path for path in out.rglob("*") if path.is_file() and not path.name.lower().endswith(RAW_EXTENSIONS)]
    rows = []
    downloaded = {p.name for p in paths}
    for record in records:
        name, category, size = file_fields(record)
        rows.append({"filename": name, "category": category, "size_bytes": size,
                     "download_status": "NOT_DOWNLOADED_RAW" if category == "RAW" else
                     "DOWNLOADED" if name in downloaded else "SKIPPED_SIZE_OR_NOT_SELECTED"})
    return paths, rows


def decode_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def iter_text_tables(path: Path) -> Iterable[tuple[str, Iterable[list[str]]]]:
    """Yield (member, rows) for supported text/spreadsheet files."""
    lower = path.name.lower()
    if lower.endswith(".zip"):
        try:
            with zipfile.ZipFile(path) as archive:
                for item in archive.infolist():
                    name = item.filename.lower()
                    if item.is_dir() or item.file_size > 250_000_000 or name.endswith(RAW_EXTENSIONS):
                        continue
                    if name.endswith(TABLE_EXTENSIONS[:-1]):
                        try:
                            raw = archive.read(item)
                            yield from iter_text_bytes(item.filename, raw)
                        except Exception:
                            continue
        except (OSError, zipfile.BadZipFile):
            return
    elif lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                def rows(sheet=ws):
                    for row in sheet.iter_rows(values_only=True):
                        yield [decode_cell(v) for v in row]
                yield f"{path.name}::{ws.title}", rows()
        except Exception:
            return
    elif lower.endswith(".xls"):
        return
    else:
        try:
            raw = path.read_bytes()
            yield from iter_text_bytes(path.name, raw)
        except OSError:
            return


def iter_text_bytes(name: str, raw: bytes) -> Iterable[tuple[str, Iterable[list[str]]]]:
    if name.lower().endswith((".gz",)):
        try:
            raw = gzip.decompress(raw)
        except OSError:
            return
    text = raw.decode("utf-8-sig", errors="replace")
    first = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(first, delimiters="\t,;")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if first.count("\t") >= first.count(",") else ","
    def rows():
        for row in csv.reader(io.StringIO(text), delimiter=delimiter):
            yield row
    yield name, rows()


def is_number(value: str) -> bool:
    value = value.strip().replace(",", "")
    if not value or value.lower() in {"nan", "na", "n/a", "null", "none", "-"}:
        return False
    try:
        return float(value) == float(value)
    except ValueError:
        return False


def clean_peptide(value: str) -> str:
    value = value.strip()
    value = re.sub(r"\[[^\]]*\]|\([^)]*(?:phosph|\+79|\+80)[^)]*\)", "", value, flags=re.I)
    value = value.replace(".", "")
    value = re.sub(r"[^A-Za-z]", "", value)
    # MaxQuant lower-case amino-acid letters encode a variable modification.
    return value.upper()


def sequence_matches_site(sequence: str, site_text: str) -> bool:
    """Require the peptide sequence to map the phosphoserine to the S15 context."""
    peptide = clean_peptide(sequence)
    if not peptide:
        return False
    context = "QEKRRQISIRGIVGV"
    # Context coordinates: the S is the 8th residue, corresponding to PYGL 15.
    explicit_site = bool(re.search(r"(?:\bS\s*15\b|\b15\s*S\b|Ser\s*15)", site_text, re.I))
    if context in peptide and explicit_site:
        return True
    raw_letters = re.sub(r"[^A-Za-z]", "", sequence)
    raw_upper = raw_letters.upper()
    offset = raw_upper.find(context)
    if offset >= 0 and raw_letters[offset + 7] == "s":
        return True
    for fragment in ("RQISIR", "QISIRGIVGV", "QISIR", "ISIRGIVGV"):
        offset = context.find(fragment)
        if offset >= 0 and fragment in peptide and context[offset + fragment.index("S")] == "S":
            # The table's site label must independently assign this sequence to S15.
            if explicit_site:
                return True
    return False


def row_quant_and_ids(header: list[str], row: list[str]) -> tuple[int, list[str], list[str]]:
    quant_indices = []
    sample_ids: set[str] = set()
    patient_ids: set[str] = set()
    for index, title in enumerate(header):
        col = title.strip()
        if not col:
            continue
        if re.search(r"intens|abundance|reporter|ratio|quant|signal|area|tmt|lfq|normalized|normalised|patient|sample", col, re.I) or re.fullmatch(r"(?:126|127[NC]?|128[NC]?|129[NC]?|130[NC]?|131[NC]?)", col, re.I):
            if index < len(row) and is_number(row[index]):
                quant_indices.append(index)
                sample_ids.add(col)
        if re.search(r"patient|subject|sample|specimen|case|tcga|\b id\b", col, re.I):
            candidates = re.findall(r"TCGA-[A-Z0-9-]+|(?:patient|pat|pt|AML|Ex|P)[-_ ]?\d{1,4}", col, re.I)
            for candidate in candidates:
                patient_ids.add(re.sub(r"\s+", "", candidate).upper())
    # In long-format exports, recover IDs only from explicitly named ID columns.
    for index, title in enumerate(header):
        if index < len(row) and re.search(r"patient|subject|sample|specimen|case|tcga|\b id\b", title, re.I):
            candidates = re.findall(r"TCGA-[A-Z0-9-]+|(?:patient|pat|pt|AML|Ex|P)[-_ ]?\d{1,4}", row[index], re.I)
            for candidate in candidates:
                patient_ids.add(re.sub(r"\s+", "", candidate).upper())
    # In long format, a single Intensity/Ratio field represents one quantified sample per row.
    return len(quant_indices), sorted(sample_ids), sorted(patient_ids)


def inspect_table(dataset: str, repo: str, path: Path, member: str, rows: Iterable[list[str]]) -> list[dict[str, Any]]:
    iterator = iter(rows)
    try:
        header = [decode_cell(v) for v in next(iterator)]
    except StopIteration:
        return []
    joined_header = " ".join(header)
    lower_header = joined_header.lower()
    protein_cols = [i for i, h in enumerate(header) if re.search(r"protein|gene|uniprot|accession|majority", h, re.I)]
    site_cols = [i for i, h in enumerate(header) if re.search(r"site|position|phosph|modification|amino", h, re.I)]
    sequence_cols = [i for i, h in enumerate(header) if re.search(r"sequence|peptide", h, re.I)]
    rows_out: list[dict[str, Any]] = []
    for row_number, cells in enumerate(iterator, start=2):
        values = [decode_cell(v) for v in cells]
        whole = " | ".join(values)
        if not re.search(r"\bPYGL\b|\bP06737\b|ENSG00000100504", whole, re.I):
            continue
        proteins = sorted({token for token in re.findall(r"P06737|ENSG00000100504|\bPYGL\b", whole, re.I)})
        sites = []
        for index in site_cols:
            if index < len(values):
                sites.append(values[index])
        site_text = " | ".join(sites + [whole])
        has_s15 = bool(re.search(r"(?:\bS\s*15\b|\b15\s*S\b|Ser\s*15)", site_text, re.I))
        is_s430 = bool(re.search(r"(?:\bS\s*430\b|\b430\s*S\b|Ser\s*430)", site_text, re.I))
        sequences = [values[i] for i in sequence_cols if i < len(values)]
        if not sequences:
            sequences = [v for v in values if re.search(r"QEKRRQISIRGIVGV|RQISIR|QISIRGIVGV", v, re.I)]
        if any(sequence_matches_site(v, site_text) for v in sequences):
            has_s15 = True
        seq = next((v for v in sequences if sequence_matches_site(v, site_text)), sequences[0] if sequences else "")
        sequence_ok = sequence_matches_site(seq, site_text) if seq else False
        numeric_n, sample_ids, patient_ids = row_quant_and_ids(header, values)
        if numeric_n == 0 and re.search(r"intens|abundance|reporter|ratio|quant|signal|area|tmt|lfq", lower_header):
            numeric_n = sum(1 for i, h in enumerate(header) if re.search(r"intens|abundance|reporter|ratio|quant|signal|area|tmt|lfq", h, re.I) and i < len(values) and is_number(values[i]))
        phospho_context = bool(site_cols or sequence_cols or re.search(r"phosph|site", lower_header, re.I))
        if not phospho_context and not has_s15:
            hit_type = "PYGL_TOTAL_PROTEIN_QUANT" if numeric_n else "PYGL_PROTEIN_IDENTIFICATION"
        elif is_s430:
            hit_type = "PYGL_OTHER_SITE_ONLY"
        elif has_s15 and sequence_ok and numeric_n:
            hit_type = "CONFIRMED_S15_PATIENT_QUANT"
        elif has_s15 and sequence_ok:
            hit_type = "CONFIRMED_S15_PSM_ONLY"
        else:
            hit_type = "PYGL_OTHER_SITE_ONLY" if not has_s15 else "AMBIGUOUS_S15_SEQUENCE"
        rows_out.append({
            "dataset": dataset, "repository_accession": repo, "filename": path.name,
            "table_member": member, "row_number": row_number, "hit_type": hit_type,
            "protein_identifiers": ",".join(proteins), "site_identifiers": site_text[:300],
            "peptide_sequence": seq[:200], "sequence_confirmed": "YES" if sequence_ok else "NO",
            "quantitative_sample_n": numeric_n, "observed_patient_n": len(patient_ids) if patient_ids else "",
            "patient_ids_recovered": "YES" if patient_ids else "NO",
            "site_localization": "S15 label + peptide/context match" if sequence_ok else "not sequence-confirmed",
            "false_positive_excluded": "YES (S430 region excluded)" if is_s430 else "NO",
        })
    return rows_out


def classify_dataset(study: dict[str, str], files: list[dict[str, Any]], file_error: str,
                     candidates: list[dict[str, Any]], scanned: bool) -> dict[str, str]:
    row = {field: "" for field in INVENTORY_FIELDS}
    row.update({"dataset": study["accession"], "disease": study["disease"], "patient_n": study["patient_n"],
                "sample_n": study["sample_n"], "tissue": study["tissue"],
                "clinical_endpoint": study["clinical_endpoint"]})
    if file_error and not files:
        row.update({"processed_phosphosite_available": "UNKNOWN", "total_proteome_available": study["total_proteome_expected"],
                    "patient_mapping_available": study["mapping_expected"], "raw_only": "UNKNOWN",
                    "status": "DATA_NOT_ACCESSIBLE"})
        return row
    names = [file_fields(x)[0].lower() for x in files]
    cats = [file_fields(x)[1] for x in files]
    processed = any(cat in {"RESULT", "SEARCH", "OTHER"} and not name.endswith(RAW_EXTENSIONS) for name, cat in zip(names, cats))
    total = any(re.search(r"protein.?groups|protein.?abundance|total.?proteome|proteome.?matrix|global.?protein", name, re.I) for name in names)
    mapping = any(re.search(r"clinical|sample.?map|metadata|patient|phenotype|survival|outcome", name, re.I) for name in names)
    row["processed_phosphosite_available"] = "YES" if processed else "NO"
    row["total_proteome_available"] = "YES" if total or study["total_proteome_expected"] == "YES" else study["total_proteome_expected"]
    row["patient_mapping_available"] = "YES (file found)" if mapping else study["mapping_expected"]
    row["raw_only"] = "YES" if not processed and any(cat == "RAW" for cat in cats) else "NO" if processed else "UNKNOWN"
    if not scanned:
        row["status"] = "NOT_SCANNED_AFTER_ELIGIBLE_HIT"
    elif any(x["hit_type"] == "CONFIRMED_S15_PATIENT_QUANT" for x in candidates):
        row["status"] = "CONFIRMED_S15_PATIENT_QUANT"
    elif any(x["hit_type"] == "CONFIRMED_S15_PSM_ONLY" for x in candidates):
        row["status"] = "CONFIRMED_S15_PSM_ONLY"
    elif any(x["hit_type"] == "PYGL_OTHER_SITE_ONLY" for x in candidates):
        row["status"] = "PYGL_OTHER_SITE_ONLY"
    else:
        row["status"] = "NO_PYGL_S15" if not file_error else "DATA_NOT_ACCESSIBLE"
    return row


def scan_dataset(pxd: str, repo: str, paths: list[Path]) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    for path in paths:
        for member, rows in iter_text_tables(path):
            found.extend(inspect_table(pxd, repo, path, member, rows))
    return found


def main() -> int:
    RESULTS.mkdir(parents=True, exist_ok=True)
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    try:
        from pridepy.download.client import Client
        client = Client()
    except Exception as exc:
        print(f"Could not initialize pridepy: {exc}", file=sys.stderr)
        return 2

    file_rows: list[dict[str, Any]] = []
    inventory_by_id: dict[str, dict[str, str]] = {}
    data_by_id: dict[str, tuple[str, str, list[dict[str, Any]], list[Path], str]] = {}
    ordered = STUDIES[:]

    # Phase 1: inventory every prespecified candidate and its file categories.
    for study in ordered:
        pxd = study["accession"]
        native, source, files, error = list_processed_files(client, pxd)
        selected = choose_downloads(files)
        for rec in files:
            name, category, size = file_fields(rec)
            file_rows.append({"dataset": pxd, "repository_accession": native,
                              "filename": name, "category": category, "size_bytes": size,
                              "download_status": "INVENTORIED"})
        data_by_id[pxd] = (native, source, files, [], error)
        inventory_by_id[pxd] = classify_dataset(study, files, error, [], scanned=False)

    # Phase 2: scan tier 1 first, then tier 2; stop at first eligible direct-S15 route.
    all_evidence: list[dict[str, Any]] = []
    eligible_hit = False
    scan_errors: dict[str, str] = {}
    for study in ordered:
        pxd = study["accession"]
        native, source, files, _, error = data_by_id[pxd]
        if eligible_hit:
            inventory_by_id[pxd] = classify_dataset(study, files, error, [], scanned=False)
            continue
        if error and not files:
            inventory_by_id[pxd] = classify_dataset(study, files, error, [], scanned=True)
            scan_errors[pxd] = error
            continue
        selected = choose_downloads(files)
        try:
            paths, outcomes = download_selected(client, native, pxd, selected)
            # Mark file-level download status while retaining all listed files.
            outcome_by_name = {x["filename"]: x["download_status"] for x in outcomes}
            for row in file_rows:
                if row["dataset"] == pxd and row["filename"] in outcome_by_name:
                    row["download_status"] = outcome_by_name[row["filename"]]
            candidates = scan_dataset(pxd, native, paths)
            all_evidence.extend(candidates)
            inventory_by_id[pxd] = classify_dataset(study, files, error, candidates, scanned=True)
            eligible_hit = any(x["hit_type"] == "CONFIRMED_S15_PATIENT_QUANT" for x in candidates)
            if eligible_hit:
                inventory_by_id[pxd]["status"] = "CONFIRMED_S15_PATIENT_QUANT"
        except Exception as exc:
            scan_errors[pxd] = f"{type(exc).__name__}: {exc}"
            inventory_by_id[pxd] = classify_dataset(study, files, f"{error}; {scan_errors[pxd]}", [], scanned=True)

    write_tsv(MANIFEST, MANIFEST_FIELDS, file_rows)
    write_tsv(EVIDENCE, EVIDENCE_FIELDS, all_evidence)
    write_tsv(INVENTORY, INVENTORY_FIELDS, [inventory_by_id[x["accession"]] for x in ordered])

    confirmed = [x for x in all_evidence if x["hit_type"] == "CONFIRMED_S15_PATIENT_QUANT"]
    pygl = [x for x in all_evidence if x["protein_identifiers"]]
    no_access = [x for x in ordered if inventory_by_id[x["accession"]]["status"] == "DATA_NOT_ACCESSIBLE"]
    lines = [
        "AUDIT_SCOPE=9 prespecified non-CPTAC datasets; no CPTAC/PDC sources queried",
        f"DATASET_INVENTORY_ROWS={len(ordered)}",
        f"DATASETS_SCANNED_FOR_PROCESSED_TABLES={sum(1 for x in ordered if inventory_by_id[x['accession']]['status'] != 'NOT_SCANNED_AFTER_ELIGIBLE_HIT')}",
        f"DATASETS_WITH_PYGL_ROWS={len(set(x['dataset'] for x in pygl))}",
        f"DATASETS_WITH_SEQUENCE_CONFIRMED_S15_PATIENT_QUANT={len(set(x['dataset'] for x in confirmed))}",
        f"DATASETS_NOT_ACCESSIBLE={','.join(x['accession'] for x in no_access) or 'NONE'}",
        f"ELIGIBLE_HIT={confirmed[0]['dataset'] if confirmed else 'NONE'}",
        "RAW_DOWNLOADS=0",
        "",
        "SCAN_ERRORS:",
    ]
    lines.extend(f"{key}: {value}" for key, value in scan_errors.items())
    STATUS.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

