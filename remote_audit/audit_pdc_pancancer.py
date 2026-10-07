#!/usr/bin/env python3
"""Remote, low-footprint audit of four CPTAC Pan-Cancer phosphoproteome files."""
import csv
import gzip
import hashlib
import io
import itertools
import json
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from collections import defaultdict
from pathlib import Path

import requests

OUT = Path("remote_results")
TMP = Path("/tmp/pdc_pygl")
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)
LOG = OUT / "remote_audit.log"
HOSTS = ("https://proteomic.datacommons.cancer.gov", "https://pdc.cancer.gov")
DELAYS = (5, 10, 20, 40, 60)
TIMEOUT = 180
FILES = (
    "Phosphoproteome_UMich_GENCODE34_v1.zip",
    "Phosphoproteome_BCM_GENCODE_v34_harmonized_v1.zip",
    "Phosphoproteome_UMich_SinaiPreprocessed_GENECODE34_v1.zip",
    "Phosphoproteome_Broad_Institute_harmonized_v1.tsv.gz",
)
COHORTS = ("BRCA", "ccRCC", "COAD", "GBM", "HGSC", "HNSCC", "LSCC", "LUAD", "PDAC", "UCEC", "MB")
SEARCH_IDS = ("PYGL", "P06737", "ENSG00000100504")
TARGET_NTERM = "MAKPLTDQEKRRQISIRGIVGV"
S15_HIT_COUNT = 0
AMBIGUOUS_S15_COUNT = 0

META_FIELDS = ("file_id", "file_name", "file_location", "file_size", "data_category",
               "data_source", "file_type", "file_format", "downloadable",
               "metadata_status", "metadata_host", "http_status", "error")
DOWNLOAD_FIELDS = ("filename", "download_success", "size", "SHA256", "HTTP_status",
                   "parse_status", "final_status", "error")
INDEX_FIELDS = ("archive", "archive_member_name", "compressed_size", "uncompressed_size",
                "format", "cohort", "rows", "columns", "index_fields", "sample_columns", "notes")
HIT_FIELDS = ("archive", "pipeline", "cohort", "source_file", "row_index", "gene",
              "protein", "site", "peptide", "modified_peptide", "row_identifier",
              "identifier_fields", "quantitative_column_count",
              "nonmissing_quantitative_value_count", "sample_ids_nonmissing",
              "classification", "canonical_site", "mapping_evidence")
S15_FIELDS = ("pipeline", "archive", "cohort", "file", "row_id", "gene", "protein",
              "reported_site", "canonical_site", "peptide", "modified_peptide",
              "mapping_confidence", "total_sample_n", "nonmissing_sample_n", "sample_ids")
CROSS_FIELDS = ("repository", "file", "line", "matched_term", "snippet", "status", "notes")


def log(message):
    stamp = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    line = f"{stamp} {message}"
    print(line, flush=True)
    with LOG.open("a", encoding="utf-8") as f:
        f.write(line + "\n")


def make_writer(path, fields):
    f = Path(path).open("w", encoding="utf-8", newline="")
    w = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n", extrasaction="ignore")
    w.writeheader()
    return f, w


def safe_error(exc):
    return re.sub(r"https?://\S+", "[URL_REDACTED]", str(exc))[:500]


def request_json(path, method="GET", payload=None):
    last_error, last_status = "request failed", ""
    for host in HOSTS:
        for attempt in range(len(DELAYS) + 1):
            try:
                response = requests.request(method, host + path, json=payload,
                                            timeout=TIMEOUT, headers={"Accept": "application/json"})
                last_status = response.status_code
                if response.status_code == 200:
                    try:
                        return response.json(), host, last_status, ""
                    except ValueError as exc:
                        last_error = f"JSON decode error: {safe_error(exc)}"
                else:
                    last_error = f"HTTP {response.status_code}"
            except requests.RequestException as exc:
                last_error = safe_error(exc)
            if attempt < len(DELAYS):
                wait = DELAYS[attempt]
                log(f"retry host={host} attempt={attempt + 1}/5 wait={wait}s status={last_status} error={last_error}")
                time.sleep(wait)
        log(f"host attempts exhausted: {host}")
    return None, "", last_status, last_error


def query_metadata(file_name):
    query = """query PancancerFileMetadata($name: String!) {
      pancancerFileMetadata(file_name: $name) {
        file_id file_name file_location file_size data_category data_source file_type file_format downloadable
      }
    }"""
    response, host, status, error = request_json(
        "/graphql", "POST", {"query": query, "variables": {"name": file_name}})
    if response is None:
        return [], "METADATA_FAILED", host, status, error
    if response.get("errors"):
        return [], "METADATA_FAILED", host, status, json.dumps(response["errors"], ensure_ascii=False)[:500]
    data = response.get("data")
    records = (data or {}).get("pancancerFileMetadata") or []
    if isinstance(records, dict):
        records = [records]
    records = [r for r in records if isinstance(r, dict) and r.get("file_name") == file_name]
    if not records:
        return [], "METADATA_FAILED", host, status, "No exact file_name record returned"
    return records, "METADATA_SUCCESS", host, status, ""


def signed_url(file_id):
    response, host, status, error = request_json(f"/pdcapi/file/signedURLFromUuid/{file_id}")
    if response is None:
        return "", host, status, error
    if response.get("error") is not False:
        return "", host, status, "signed URL response error flag was not false"
    url = response.get("data")
    if isinstance(url, dict):
        url = url.get("url") or url.get("signedUrl") or url.get("signed_url")
    if not isinstance(url, str) or not url.startswith("https://"):
        return "", host, status, "response contained no HTTPS data URL"
    return url, host, status, ""


def download(file_name, file_id):
    url, _, status, error = signed_url(file_id)
    if not url:
        return None, False, 0, "", status, error
    target = TMP / (str(file_id) + "_" + Path(file_name).name)
    last_error, last_status = error, status
    for attempt in range(len(DELAYS) + 1):
        if attempt:
            url, _, last_status, last_error = signed_url(file_id)
            if not url:
                target.unlink(missing_ok=True)
                if attempt < len(DELAYS):
                    log(f"retry signed URL file={file_name} attempt={attempt}/5 error={last_error}")
                    time.sleep(DELAYS[attempt - 1])
                continue
        digest, size = hashlib.sha256(), 0
        try:
            with requests.get(url, stream=True, timeout=TIMEOUT) as response:
                last_status = response.status_code
                if 200 <= last_status < 300:
                    with target.open("wb") as out:
                        for chunk in response.iter_content(1024 * 1024):
                            if chunk:
                                out.write(chunk)
                                digest.update(chunk)
                                size += len(chunk)
                    return target, True, size, digest.hexdigest(), last_status, ""
                last_error = f"download HTTP {last_status}"
        except (requests.RequestException, OSError) as exc:
            last_error = safe_error(exc)
        target.unlink(missing_ok=True)
        if attempt < len(DELAYS):
            log(f"retry download file={file_name} attempt={attempt + 1}/5 wait={DELAYS[attempt]}s status={last_status} error={last_error}")
            time.sleep(DELAYS[attempt])
    return None, False, 0, "", last_status, last_error


def norm(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def missing(value):
    return value is None or (isinstance(value, str) and value.strip().lower() in
                             ("", "na", "n/a", "nan", "null", "none", "-", "."))


def numeric(value):
    if missing(value):
        return False
    try:
        float(str(value).replace(",", ""))
        return True
    except (TypeError, ValueError):
        return False


def choose_header(rows):
    terms = ("gene", "symbol", "protein", "accession", "site", "phospho",
             "peptide", "sequence", "identifier", "feature", "cohort", "cancer")
    scores = [sum(term in " ".join(norm(x) for x in row) for term in terms) for row in rows]
    return max(range(len(scores)), key=scores.__getitem__) if scores else 0


def read_delimited(stream, name):
    text = io.TextIOWrapper(stream, encoding="utf-8-sig", errors="replace", newline="")
    first = text.readline()
    if not first:
        return [], iter(()), 0, "empty text file"
    delim = "\t" if first.count("\t") >= max(first.count(","), first.count(";")) else ("," if first.count(",") >= first.count(";") else ";")
    reader = csv.reader(itertools.chain((first,), text), delimiter=delim)
    preview = list(itertools.islice(reader, 10))
    if not preview:
        return [], iter(()), 0, "empty text file"
    hidx = choose_header(preview)
    header = [str(x or "").strip() for x in preview[hidx]]
    remaining = itertools.chain(preview[hidx + 1:], reader)
    note = f"delimiter={repr(delim)}; header_row={hidx+1}; preview={json.dumps(preview[:3], ensure_ascii=False)[:1000]}"
    return header, remaining, hidx + 1, note


META_TOKENS = ("gene", "symbol", "protein", "accession", "uniprot", "ensembl", "site",
               "phospho", "peptide", "sequence", "localization", "probability",
               "identifier", "feature", "description", "cohort", "cancer", "case",
               "patient", "sampleid", "run", "condition", "group", "database",
               "rowid", "index", "position", "maxpepprob", "referenceintensity")


def split_columns(header, first_row):
    metadata, samples = [], []
    for i, name in enumerate(header):
        n = norm(name)
        if any(t in n for t in META_TOKENS):
            metadata.append(i)
        elif i < len(first_row) and (numeric(first_row[i]) or missing(first_row[i])):
            samples.append(i)
        else:
            metadata.append(i)
    return metadata, samples


def field(header, row, tokens, exclude=()):
    excludes = tuple(norm(x) for x in exclude)
    for token in tokens:
        needle = norm(token)
        for i, name in enumerate(header):
            n = norm(name)
            matches = n in ("id", "rowid", "featureid", "proteinid") if needle == "id" else needle in n
            if (matches and not any(x in n for x in excludes)
                    and i < len(row) and not missing(row[i])):
                return str(row[i]).strip()
    return ""


def metadata_values(header, row):
    out = {}
    for i, name in enumerate(header):
        if any(t in norm(name) for t in META_TOKENS):
            out[str(name)] = row[i] if i < len(row) else ""
    return out


def site_positions(text):
    patterns = (
        r"(?i)(?:^|[^a-z])(?:p)?s(?:er)?\s*[_-]?\s*0*(\d+)(?=$|[^0-9])",
        r"(?i)(?:^|[^0-9])(\d+)\s*(?:p)?s(?:er)?(?=$|[^a-z])",
        r"(?i)p06737[_:.-]+s[_-]?0*(\d+)",
        r"(?i)pygl[_:.-]+s[_-]?0*(\d+)",
    )
    found = set()
    for pattern in patterns:
        found.update(int(m.group(1)) for m in re.finditer(pattern, str(text or "")))
    for match in re.finditer(r"(?i)\|\d+_(\d+)_(\d+)(?=_|$)", str(text or "")):
        found.update(int(value) for value in match.groups())
    return found


def peptide_sequence(value):
    """Strip modification markup and return AA sequence plus explicit phospho offsets."""
    s = str(value or "").strip()
    if re.match(r"^[A-Z-]\.[A-Za-z\[\(].*\.[A-Z-]$", s):
        s = s[2:-2]
    seq, mods, i = [], set(), 0
    while i < len(s):
        ch = s[i]
        if ch == "[":
            end = s.find("]", i + 1)
            if end >= 0:
                token = s[i + 1:end]
                aa = re.search(r"(?i)([STY])", token)
                if aa:
                    seq.append(aa.group(1).upper())
                    if re.search(r"(?i)(ph|phospho|\+79\.966|79\.966|p[sty])", token):
                        mods.add(len(seq) - 1)
                i = end + 1
                continue
        if ch.isalpha() and ch.upper() in "ACDEFGHIKLMNPQRSTVWY":
            seq.append(ch.upper())
            idx = len(seq) - 1
            if ch in "sty":
                mods.add(idx)
            end = i + 1
            if end < len(s) and s[end] in "([":
                close = ")" if s[end] == "(" else "]"
                stop = s.find(close, end + 1)
                if stop >= 0:
                    token = s[end + 1:stop]
                    if ch.upper() in token.upper() or re.search(r"(?i)(ph|phospho|\+79\.966|79\.966)", token):
                        mods.add(idx)
                    i = stop + 1
                    continue
            i += 1
            continue
        i += 1
    return "".join(seq), mods


def map_peptide(value, sequence):
    peptide, mods = peptide_sequence(value)
    starts, start = [], sequence.find(peptide) if peptide else -1
    while start >= 0:
        starts.append(start)
        start = sequence.find(peptide, start + 1)
    return peptide, mods, starts


def classify(site, peptide, modified, sequence):
    raw_values = modified or peptide
    alternatives = [x.strip() for x in str(raw_values or "").split(";") if x.strip()] or [""]
    outcomes = []
    for raw in alternatives:
        result = _classify_single(site, raw, sequence)
        if result[0] == "CONFIRMED_S15":
            return result
        outcomes.append(result)
    for priority in ("FALSE_POSITIVE_SITE_LABEL", "OTHER_PYGL_SITE", "AMBIGUOUS"):
        for result in outcomes:
            if result[0] == priority:
                return result
    return "AMBIGUOUS", "", "no peptide alternative could be classified"


def _classify_single(site, raw, sequence):
    site_text = str(site or "")
    positions = site_positions(site_text)
    aa, mods = peptide_sequence(raw)
    if 15 in positions and re.search(r"(?i)(?:RMSLIEEEG|MSLIEEEG)", aa):
        return "FALSE_POSITIVE_SITE_LABEL", "S430", "fixed C-terminal RMSLIEEEG peptide control"
    seq, mods, starts = map_peptide(raw, sequence)
    if len(starts) != 1:
        return "AMBIGUOUS", "", "peptide does not map uniquely to canonical P06737"
    start = starts[0]
    mod_sites = {start + i + 1 for i in mods}
    if 15 in mod_sites and start <= 14 < start + len(seq) and seq[14 - start] == "S":
        return "CONFIRMED_S15", "S15", "unique canonical peptide alignment and phospho-Ser at residue 15"
    if 15 in positions:
        if start <= 14 < start + len(seq) and seq[14 - start] == "S":
            if re.search(r"(?i)(?<![A-Z0-9])S0*15(?!\d)", site_text):
                evidence = "reported S15 plus unique peptide alignment to canonical Ser15"
            else:
                evidence = "numeric site coordinate 15 plus unique peptide alignment identifies canonical Ser15"
            return "CONFIRMED_S15", "S15", evidence
        return "FALSE_POSITIVE_SITE_LABEL", "", "reported S15 conflicts with peptide alignment"
    mapped = {start + i + 1 for i in range(len(seq))}
    if mod_sites:
        return "OTHER_PYGL_SITE", ";".join(str(x) for x in sorted(mod_sites)), "explicit modification maps to a different canonical position"
    if positions and any(p in mapped for p in positions):
        return "OTHER_PYGL_SITE", ";".join(str(x) for x in sorted(positions)), "reported coordinate agrees with peptide alignment"
    return "AMBIGUOUS", "", "site label or peptide alone is insufficient for mapping"


def could_be_ambiguous_s15(site, peptide, modified, sequence):
    """Flag unresolved evidence that could still map a phosphopeptide to canonical S15."""
    if 15 in site_positions(site):
        return True
    raw_values = modified or peptide
    for raw in str(raw_values or "").split(";"):
        _, mods, starts = map_peptide(raw.strip(), sequence)
        if any(start + offset + 1 == 15 for start in starts for offset in mods):
            return True
    return False


def get_reference():
    seq, ensp = "", set()
    try:
        r = requests.get("https://rest.uniprot.org/uniprotkb/P06737.fasta", timeout=TIMEOUT)
        if r.status_code == 200:
            seq = "".join(x.strip() for x in r.text.splitlines() if x and not x.startswith(">"))
        else:
            log(f"UniProt FASTA HTTP {r.status_code}")
    except requests.RequestException as exc:
        log(f"UniProt FASTA error: {safe_error(exc)}")
    try:
        r = requests.get("https://rest.uniprot.org/uniprotkb/P06737.json", timeout=TIMEOUT)
        if r.status_code == 200:
            for item in r.json().get("uniProtKBCrossReferences", []):
                if item.get("database") == "Ensembl" and str(item.get("id", "")).startswith("ENSP"):
                    ensp.add(item["id"])
    except (requests.RequestException, ValueError) as exc:
        log(f"UniProt ENSP lookup error: {safe_error(exc)}")
    if not seq:
        seq = TARGET_NTERM
        log("Using target N-terminal sequence fallback; non-S15 site mapping may be AMBIGUOUS")
    return seq, ensp


def cohort(text):
    s = str(text or "")
    for code in COHORTS:
        if code == "ccRCC" and re.search(r"(?i)cc[-_ ]?rcc|clear.?cell", s):
            return "ccRCC"
        if code == "HGSC" and re.search(r"(?i)HGSC|OVARIAN", s):
            return "HGSC/OV"
        if re.search(rf"(?i)(?<![A-Z0-9]){re.escape(code)}(?![A-Z0-9])", s):
            return code
    return "UNKNOWN"


def scan_rows(header, rows, archive, source, pipeline, sequence, ensp, hit_writer, s15_writer, coverage):
    global S15_HIT_COUNT, AMBIGUOUS_S15_COUNT
    iterator = iter(rows)
    first = next(iterator, None)
    if first is None:
        return 0, 0, [], [], "header only"
    first = ["" if x is None else str(x) for x in first]
    metadata_idx, sample_idx = split_columns(header, first)
    sample_names = [header[i] for i in sample_idx]
    index_names = [header[i] for i in metadata_idx]
    total, matched = 0, 0
    for values in itertools.chain((first,), iterator):
        values = ["" if x is None else str(x) for x in values]
        if not values or all(missing(x) for x in values):
            continue
        total += 1
        joined = " | ".join(values).upper()
        target_identity = (
            re.search(r"(?<![A-Z0-9])ENSG00000100504(?:\.\d+)?(?![A-Z0-9])", joined)
            or re.search(r"(?<![A-Z0-9])P06737(?![A-Z0-9])", joined)
            or re.search(r"(?<![A-Z0-9])PYGL(?![A-Z0-9])", joined)
            or any(re.search(rf"(?<![A-Z0-9]){re.escape(x.upper())}(?![A-Z0-9])", joined) for x in ensp)
        )
        if not target_identity:
            continue
        matched += 1
        gene = field(header, values, ("gene", "symbol"))
        row_id = field(header, values, ("rowid", "identifier", "feature", "id", "index", "description"))
        protein = field(header, values, ("uniprot", "accession", "proteinid", "protein"))
        peptide = field(header, values, ("peptide",), ("modified",)) or field(header, values, ("sequence",))
        modified = field(header, values, ("modifiedpeptide", "modifiedsequence", "modpeptide"))
        identifier_values = [values[i] for i in metadata_idx if i < len(values)
                             and any(term in values[i].upper() for term in SEARCH_IDS + tuple(ensp))]
        identifier_evidence = " | ".join(identifier_values)
        row_id = row_id or identifier_evidence or " | ".join(x for x in (gene, protein, peptide) if x)
        if not gene and ("P06737" in joined or "ENSG00000100504" in joined
                         or any(x.upper() in joined for x in ensp)):
            gene = "PYGL"
        if not protein:
            protein = "P06737" if "P06737" in joined else next(
                (x for x in ensp if x.upper() in joined), "")
        site = field(header, values, ("phosphosite", "site", "position"),
                     ("probability", "localization", "confidence", "score")) or row_id or identifier_evidence
        row_cohort = field(header, values, ("cohort", "cancer", "disease"))
        explicit_cohort = cohort(row_cohort) if row_cohort else "UNKNOWN"
        source_cohort = cohort(source)
        sample_cohorts = {}
        for i in sample_idx:
            inferred = explicit_cohort if explicit_cohort != "UNKNOWN" else cohort(header[i])
            sample_cohorts[i] = inferred if inferred != "UNKNOWN" else source_cohort
        known_cohorts = sorted({value for value in sample_cohorts.values() if value != "UNKNOWN"})
        cohort_name = explicit_cohort if explicit_cohort != "UNKNOWN" else (
            known_cohorts[0] if len(known_cohorts) == 1 else
            ("PAN_CANCER" if len(known_cohorts) > 1 else source_cohort))
        if not cohort_name:
            cohort_name = "UNKNOWN"
        nonmissing = [header[i] for i in sample_idx if i < len(values) and numeric(values[i])]
        call, canonical, why = classify(site, peptide, modified, sequence)
        meta = {header[i]: values[i] for i in metadata_idx if i < len(values)}
        hit_writer.writerow({
            "archive": archive, "pipeline": pipeline, "cohort": cohort_name, "source_file": source,
            "row_index": total, "gene": gene, "protein": protein, "site": site, "peptide": peptide,
            "modified_peptide": modified, "row_identifier": row_id,
            "identifier_fields": json.dumps(meta, ensure_ascii=False, separators=(",", ":")),
            "quantitative_column_count": len(sample_idx),
            "nonmissing_quantitative_value_count": len(nonmissing),
            "sample_ids_nonmissing": ";".join(nonmissing), "classification": call,
            "canonical_site": canonical, "mapping_evidence": why})
        if call == "AMBIGUOUS" and could_be_ambiguous_s15(site, peptide, modified, sequence):
            AMBIGUOUS_S15_COUNT += 1
        if call == "CONFIRMED_S15":
            S15_HIT_COUNT += 1
            s15_writer.writerow({
                "pipeline": pipeline, "archive": archive, "cohort": cohort_name, "file": source,
                "row_id": row_id, "gene": gene, "protein": protein, "reported_site": site,
                "canonical_site": "S15", "peptide": peptide, "modified_peptide": modified,
                "mapping_confidence": "HIGH: unique P06737 peptide mapping to Ser15",
                "total_sample_n": len(sample_idx), "nonmissing_sample_n": len(nonmissing),
                "sample_ids": ";".join(nonmissing)})
        if canonical or site:
            site_key = canonical or site or "UNKNOWN"
            for i in sample_idx:
                ca = sample_cohorts.get(i, "UNKNOWN")
                entry = coverage[(pipeline, ca, site_key)]
                entry["total"].add(header[i])
                if i < len(values) and numeric(values[i]):
                    entry["nonmissing"].add(header[i])
                entry["files"].add(source)
    return total, matched, index_names, sample_names, "streaming scan; first rows recorded in notes"


def pipeline(file_name):
    n = file_name.lower()
    if "bcm" in n: return "BCM"
    if "sinai" in n: return "UMich_Sinai"
    if "broad" in n: return "Broad"
    return "UMich"


def scan_zip(path, name, seq, ensp, hit_writer, s15_writer, idx_writer, coverage):
    total, hits, parsed, errors = 0, 0, 0, 0
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            for item in members:
                lower = item.filename.lower()
                if item.is_dir():
                    continue
                if lower.startswith("__macosx/") or Path(lower).name.startswith("._"):
                    idx_writer.writerow({"archive": name, "archive_member_name": item.filename,
                        "compressed_size": item.compress_size, "uncompressed_size": item.file_size,
                        "format": Path(lower).suffix.lstrip("."), "cohort": cohort(item.filename),
                        "rows": "SKIPPED_METADATA_RESOURCE", "columns": "", "index_fields": "",
                        "sample_columns": "", "notes": "AppleDouble/macOS metadata resource; excluded from matrix parsing"})
                    continue
                supported = lower.endswith((".tsv", ".csv", ".txt", ".tsv.gz", ".csv.gz", ".xlsx"))
                if any(x in lower for x in (".mzml", ".raw", "/raw", "spectra")):
                    idx_writer.writerow({"archive": name, "archive_member_name": item.filename,
                        "compressed_size": item.compress_size, "uncompressed_size": item.file_size,
                        "format": Path(lower).suffix.lstrip("."), "cohort": cohort(item.filename),
                        "rows": "SKIPPED", "columns": "", "index_fields": "", "sample_columns": "",
                        "notes": "raw/spectra member excluded"})
                    continue
                if not supported:
                    idx_writer.writerow({"archive": name, "archive_member_name": item.filename,
                        "compressed_size": item.compress_size, "uncompressed_size": item.file_size,
                        "format": Path(lower).suffix.lstrip("."), "cohort": cohort(item.filename),
                        "rows": "NOT_A_SUPPORTED_MATRIX", "columns": "", "index_fields": "",
                        "sample_columns": "", "notes": "member inventoried; not parsed"})
                    continue
                try:
                    with archive.open(item) as raw:
                        if lower.endswith(".xlsx"):
                            from openpyxl import load_workbook
                            temp = tempfile.SpooledTemporaryFile(max_size=8*1024*1024)
                            shutil.copyfileobj(raw, temp); temp.seek(0)
                            book = load_workbook(temp, read_only=True, data_only=True)
                            try:
                                for sheet in book.worksheets:
                                    iterator = sheet.iter_rows(values_only=True)
                                    preview = list(itertools.islice(iterator, 10))
                                    if not preview: continue
                                    h = choose_header(preview)
                                    header = [str(x or "").strip() for x in preview[h]]
                                    nrows, nhits, inds, samples, notes = scan_rows(
                                        header, itertools.chain(preview[h+1:], iterator), name,
                                        item.filename + "#" + sheet.title, pipeline(name), seq, ensp,
                                        hit_writer, s15_writer, coverage)
                                    total += nrows; hits += nhits; parsed += 1
                                    idx_writer.writerow({"archive": name, "archive_member_name": item.filename + "#" + sheet.title,
                                        "compressed_size": item.compress_size, "uncompressed_size": item.file_size,
                                        "format": "xlsx", "cohort": cohort(item.filename + sheet.title),
                                        "rows": nrows, "columns": len(header), "index_fields": ";".join(inds),
                                        "sample_columns": ";".join(samples), "notes": json.dumps(preview[:3], default=str, ensure_ascii=False)[:1000]})
                            finally:
                                book.close(); temp.close()
                            continue
                        stream = gzip.GzipFile(fileobj=raw, mode="rb") if lower.endswith(".gz") else raw
                        header, rows, level_count, preview = read_delimited(stream, item.filename)
                        if not header: raise ValueError(preview)
                        nrows, nhits, inds, samples, _ = scan_rows(
                            header, rows, name, item.filename, pipeline(name), seq, ensp,
                            hit_writer, s15_writer, coverage)
                        total += nrows; hits += nhits; parsed += 1
                        idx_writer.writerow({"archive": name, "archive_member_name": item.filename,
                            "compressed_size": item.compress_size, "uncompressed_size": item.file_size,
                            "format": "gzip+tsv" if lower.endswith(".gz") else Path(lower).suffix.lstrip("."),
                            "cohort": cohort(item.filename), "rows": nrows, "columns": len(header),
                            "index_fields": ";".join(inds), "sample_columns": ";".join(samples), "notes": preview})
                except Exception as exc:
                    errors += 1
                    log(f"member parse failed archive={name} member={item.filename} error={safe_error(exc)}")
                    idx_writer.writerow({"archive": name, "archive_member_name": item.filename,
                        "compressed_size": item.compress_size, "uncompressed_size": item.file_size,
                        "format": Path(lower).suffix.lstrip("."), "cohort": cohort(item.filename),
                        "rows": "PARSE_FAILED", "columns": "", "index_fields": "", "sample_columns": "",
                        "notes": safe_error(exc)})
    except Exception as exc:
        return False, total, hits, parsed, errors, safe_error(exc)
    return parsed > 0 and errors == 0, total, hits, parsed, errors, ""


def scan_gzip(path, name, seq, ensp, hit_writer, s15_writer, idx_writer, coverage):
    try:
        with gzip.open(path, "rb") as stream:
            header, rows, _, preview = read_delimited(stream, path.name)
            if not header: return False, 0, 0, 0, 1, preview
            nrows, nhits, inds, samples, _ = scan_rows(
                header, rows, name, path.name, pipeline(name), seq, ensp,
                hit_writer, s15_writer, coverage)
            idx_writer.writerow({"archive": name, "archive_member_name": path.name,
                "compressed_size": path.stat().st_size, "uncompressed_size": "STREAMED",
                "format": "gzip+tsv", "cohort": cohort(path.name), "rows": nrows,
                "columns": len(header), "index_fields": ";".join(inds),
                "sample_columns": ";".join(samples), "notes": preview})
        return True, nrows, nhits, 1, 0, ""
    except Exception as exc:
        return False, 0, 0, 0, 1, safe_error(exc)


def scan_cophee():
    repo = TMP / "CoPheeMap"
    if repo.exists(): shutil.rmtree(repo)
    handle, writer = make_writer(OUT / "CoPheeMap_PYGL_crosscheck.tsv", CROSS_FIELDS)
    clone = subprocess.run(["git", "clone", "--depth", "1", "https://github.com/bzhanglab/CoPheeMap.git", str(repo)],
                           capture_output=True, text=True)
    if clone.returncode:
        writer.writerow({"repository": "bzhanglab/CoPheeMap", "status": "CLONE_FAILED",
                         "notes": safe_error(clone.stderr or clone.stdout)})
        handle.close(); log("CoPheeMap clone failed"); return
    terms = ("PYGL", "P06737", "ENSG00000100504", "QEKRRQISIRGIVGV", "RQISIR")
    count = 0
    for path in repo.rglob("*"):
        if not path.is_file() or ".git" in path.parts or path.stat().st_size > 25*1024*1024: continue
        for line_no, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
            for term in terms:
                if term.lower() in line.lower():
                    writer.writerow({"repository": "bzhanglab/CoPheeMap",
                        "file": str(path.relative_to(repo)), "line": line_no, "matched_term": term,
                        "snippet": line[:1000], "status": "MATCH",
                        "notes": "filtered derived-data cross-check only"})
                    count += 1
    if not count:
        writer.writerow({"repository": "bzhanglab/CoPheeMap", "status": "NO_MATCH_IN_REPOSITORY",
                         "notes": "filtered network absence does not establish matrix absence"})
    handle.close(); shutil.rmtree(repo, ignore_errors=True)
    log(f"CoPheeMap read-only search complete matches={count}")


def main():
    for name in ("pancancer_file_metadata.tsv", "download_audit.tsv", "archive_inventory.tsv",
                 "PYGL_ALL_HITS.tsv", "PYGL_S15_CONFIRMED.tsv", "PYGL_SITE_COHORT_COVERAGE.tsv",
                 "PanCancer_Pipeline_Reconciliation.tsv", "FINAL_STATUS.txt", "remote_audit.log"):
        (OUT / name).unlink(missing_ok=True)
    mh, mw = make_writer(OUT / "pancancer_file_metadata.tsv", META_FIELDS)
    dh, dw = make_writer(OUT / "download_audit.tsv", DOWNLOAD_FIELDS)
    ih, iw = make_writer(OUT / "archive_inventory.tsv", INDEX_FIELDS)
    hh, hw = make_writer(OUT / "PYGL_ALL_HITS.tsv", HIT_FIELDS)
    sh, sw = make_writer(OUT / "PYGL_S15_CONFIRMED.tsv", S15_FIELDS)
    sequence, ensp = get_reference()
    log(f"P06737 reference length={len(sequence)}; ENSP crossrefs={len(ensp)}")
    metadata, metadata_count = {}, 0
    # Attempt metadata for all four before processing any archive.
    for name in FILES:
        rows, status, host, http, error = query_metadata(name)
        metadata[name] = rows
        metadata_count += status == "METADATA_SUCCESS"
        if rows:
            for row in rows:
                out = {k: row.get(k, "") for k in META_FIELDS if k in row}
                out.update(metadata_status=status, metadata_host=host, http_status=http, error="")
                mw.writerow(out)
        else:
            mw.writerow({"file_name": name, "metadata_status": status, "metadata_host": host,
                         "http_status": http or "", "error": error})
        log(f"metadata {status} file={name} http={http} error={error}")
    mh.close()
    coverage = defaultdict(lambda: {"total": set(), "nonmissing": set(), "files": set()})
    statuses, downloaded, parsed, found_s15 = [], 0, 0, False
    pygl_pipelines = set()
    for name in FILES:
        rec = next((r for r in metadata.get(name, []) if r.get("file_name") == name), None)
        if rec is None:
            status = "METADATA_FAILED"
            dw.writerow({"filename": name, "download_success": "NO", "parse_status": "NOT_RUN", "final_status": status, "error": "metadata failed"})
            statuses.append((name, status, 0, 0, "metadata failed")); continue
        if str(rec.get("downloadable", "")).lower() not in ("yes", "true", "1"):
            status = "DOWNLOAD_FAILED"
            dw.writerow({"filename": name, "download_success": "NO", "parse_status": "NOT_RUN", "final_status": status, "error": "downloadable is not Yes"})
            statuses.append((name, status, 0, 0, "downloadable is not Yes")); continue
        path, ok, size, sha, http, error = download(name, rec.get("file_id"))
        if not ok:
            status = "DOWNLOAD_FAILED"
            dw.writerow({"filename": name, "download_success": "NO", "size": size, "SHA256": sha, "HTTP_status": http or "", "parse_status": "NOT_RUN", "final_status": status, "error": error})
            statuses.append((name, status, 0, 0, error)); log(f"download failed file={name} HTTP={http} error={error}"); continue
        downloaded += 1
        log(f"download complete file={name} bytes={size} sha256={sha} HTTP={http}")
        if name.lower().endswith(".zip"):
            okparse, nrows, nhits, nfiles, nerrors, error = scan_zip(path, name, sequence, ensp, hw, sw, iw, coverage)
        else:
            okparse, nrows, nhits, nfiles, nerrors, error = scan_gzip(path, name, sequence, ensp, hw, sw, iw, coverage)
        hh.flush(); sh.flush(); ih.flush()
        if okparse:
            parsed += 1
            if nhits: pygl_pipelines.add(pipeline(name))
        if S15_HIT_COUNT:
            status = "AUDITED_S15_FOUND"
        elif not okparse:
            status = "PARSE_FAILED"
        else:
            status = "AUDITED_PYGL_NO_S15" if nhits else "AUDITED_NO_PYGL"
        dw.writerow({"filename": name, "download_success": "YES", "size": size, "SHA256": sha,
                     "HTTP_status": http or "", "parse_status": "PARSED" if okparse else "PARSE_FAILED",
                     "final_status": status, "error": error})
        statuses.append((name, status, size, nrows, error))
        log(f"scan {status} file={name} rows={nrows} PYGL_hits={nhits} members={nfiles} parse_failures={nerrors}")
        path.unlink(missing_ok=True)
        if status == "AUDITED_S15_FOUND":
            found_s15 = True
            for remaining in FILES[FILES.index(name)+1:]:
                dw.writerow({"filename": remaining, "download_success": "NOT_ATTEMPTED",
                             "parse_status": "NOT_RUN", "final_status": "SKIPPED_AFTER_S15_FOUND",
                             "error": "metadata queried; scan stopped after confirmed S15"})
                statuses.append((remaining, "SKIPPED_AFTER_S15_FOUND", 0, 0, ""))
            break
    dh.close(); ih.close(); hh.close(); sh.close()
    # Unique sample IDs are counted per cohort/site; replicate columns are not patients.
    ch, cw = make_writer(OUT / "PYGL_SITE_COHORT_COVERAGE.tsv",
                         ("pipeline", "cohort", "PYGL site", "total_samples", "nonmissing_samples", "sample_ids_nonmissing", "source_files"))
    for (pipe, ca, site), item in sorted(coverage.items()):
        if site == "UNKNOWN": continue
        cw.writerow({"pipeline": pipe, "cohort": ca, "PYGL site": site,
                     "total_samples": len(item["total"]), "nonmissing_samples": len(item["nonmissing"]),
                     "sample_ids_nonmissing": ";".join(sorted(item["nonmissing"])),
                     "source_files": ";".join(sorted(item["files"]))})
    ch.close()
    with (OUT / "PYGL_ALL_HITS.tsv").open(encoding="utf-8") as f:
        hit_rows = list(csv.DictReader(f, delimiter="\t"))
    groups = defaultdict(dict)
    for row in hit_rows:
        site = row.get("canonical_site") or row.get("site") or "UNKNOWN"
        groups[(row.get("cohort", "UNKNOWN"), site)].setdefault(row.get("pipeline", ""), []).append(row)
    rh, rw = make_writer(OUT / "PanCancer_Pipeline_Reconciliation.tsv",
                         ("cohort", "PYGL site", "BCM", "UMich_Sinai", "other_matrix", "agreement", "notes"))
    for (ca, site), by_pipe in sorted(groups.items()):
        def summary(pipe):
            rows = by_pipe.get(pipe, [])
            if not rows: return "NOT_DETECTED_OR_NOT_AUDITED"
            return f"PRESENT; rows={len(rows)}"
        states = [bool(by_pipe.get(p)) for p in ("BCM", "UMich_Sinai", "UMich", "Broad")]
        agreement = "AGREE_PRESENT" if all(states) else ("DISAGREE" if any(states) else "INCOMPLETE_PIPELINE_COVERAGE")
        other = " | ".join(f"{p}={summary(p)}" for p in ("UMich", "Broad") if by_pipe.get(p)) or "NOT_DETECTED_OR_NOT_AUDITED"
        rw.writerow({"cohort": ca, "PYGL site": site, "BCM": summary("BCM"),
                     "UMich_Sinai": summary("UMich_Sinai"), "other_matrix": other,
                     "agreement": agreement, "notes": "Sequence classification and full pipeline coverage must be considered before biological interpretation."})
    rh.close()
    complete = len(statuses) == 4 and all(s in ("AUDITED_NO_PYGL", "AUDITED_PYGL_NO_S15", "AUDITED_S15_FOUND") for _, s, _, _, _ in statuses)
    if found_s15:
        direct, conclusion, next_step = "YES", "DIRECT_S15_FOUND=YES", "CLINICAL_MAPPING"
    elif complete and AMBIGUOUS_S15_COUNT == 0:
        direct, conclusion, next_step = "NO", "NO_CANONICAL_S15_IN_PANCANCER_PROCESSED_MATRICES", "INDIVIDUAL_PDC_STUDY_AUDIT"
    else:
        direct, conclusion, next_step = "INDETERMINATE", "INCOMPLETE_REMOTE_AUDIT", "TECHNICAL_BLOCK_REMAINS"
    site_names = sorted({r.get("canonical_site") or r.get("site", "") for r in hit_rows if r.get("canonical_site") or r.get("site")})
    status_text = [f"DIRECT_S15_FOUND={direct}", f"CONCLUSION={conclusion}",
                   f"METADATA_SUCCESS={metadata_count}/4", f"DOWNLOAD_SUCCESS={downloaded}/4",
                   f"PARSE_SUCCESS={parsed}/4", f"PIPELINES_WITH_PYGL={len(pygl_pipelines)}",
                   f"AMBIGUOUS_S15_CANDIDATES={AMBIGUOUS_S15_COUNT}",
                   f"PYGL_SITE_LABELS={';'.join(site_names) or 'NONE_OBSERVED'}",
                   f"NEXT_STEP={next_step}", "FILE_STATUSES:"]
    status_text += [f"{name}\t{status}\tsize={size}\trows={rows}\t{error}" for name, status, size, rows, error in statuses]
    (OUT / "FINAL_STATUS.txt").write_text("\n".join(status_text) + "\n", encoding="utf-8")
    log(f"FINAL metadata={metadata_count}/4 download={downloaded}/4 parse={parsed}/4 pipelines_with_PYGL={len(pygl_pipelines)} S15={direct} ambiguous_S15={AMBIGUOUS_S15_COUNT} next={next_step}")
    scan_cophee()
    log("Audit finished. Archives are temporary; artifact directory contains compact outputs only.")
    print("\nREMOTE AUDIT SUMMARY")
    print(f"A. metadata success: {metadata_count}/4")
    print(f"B. download success: {downloaded}/4")
    print(f"C. parse success: {parsed}/4")
    print(f"D. pipelines detecting PYGL: {len(pygl_pipelines)}")
    print(f"E. PYGL sites: {'; '.join(site_names) or 'none observed'}")
    print(f"F. canonical S15: {direct}")
    print(f"G. confirmed S15 cohort counts: see {OUT / 'PYGL_S15_CONFIRMED.tsv'}")
    print(f"H. 4/4 complete: {'YES' if complete else 'NO'}")
    print(f"I. next step: {next_step}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        log(f"FATAL {type(exc).__name__}: {safe_error(exc)}")
        (OUT / "FINAL_STATUS.txt").write_text(
            f"DIRECT_S15_FOUND=INDETERMINATE\nCONCLUSION=REMOTE_AUDIT_FATAL\nERROR={safe_error(exc)}\nNEXT_STEP=TECHNICAL_BLOCK_REMAINS\n",
            encoding="utf-8")
        raise

